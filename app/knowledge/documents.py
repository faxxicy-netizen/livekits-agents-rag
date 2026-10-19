"""Persistent, editable knowledge documents. Files are only used for first-run demo import."""

import hashlib
import io
import re
import sqlite3
import zipfile
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path

from app.config import Settings
from app.knowledge.repository import DEFAULT_KNOWLEDGE_BASE, KnowledgeRepository
from app.models import Citation

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED = {".md", ".txt", ".pdf", ".docx"}


class DocumentError(ValueError):
    pass


def validate_filename(filename: str):
    if (
        not filename
        or len(filename) > 160
        or re.search(r"[\\/:\x00-\x1f]", filename)
        or filename in {".", ".."}
        or Path(filename).suffix.lower() not in SUPPORTED
    ):
        raise DocumentError("仅支持不含路径的 TXT、Markdown、PDF、DOCX 文件名。")


def extract_text(filename: str, data: bytes) -> str:
    validate_filename(filename)
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("单份文档不能超过 10 MB。")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".md", ".txt"}:
            try:
                content = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                content = data.decode("gb18030")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            pdf = PdfReader(io.BytesIO(data))
            if pdf.is_encrypted or len(pdf.pages) > 100:
                raise DocumentError("请上传未加密且不超过 100 页的 PDF。")
            content = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
        else:
            from docx import Document

            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise DocumentError("DOCX 解压后过大，请拆分文档。")
            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            content = "\n".join(parts)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("无法读取此文件，请检查编码或文件格式。") from exc
    content = content.replace("\x00", "").strip()
    if len(content) < 10:
        raise DocumentError("未提取到足够文字；扫描 PDF 请先做 OCR，或上传文字版文件。")
    if len(content) > 1_000_000:
        raise DocumentError("提取文字超过 100 万字，请拆分文档。")
    return content


def build_document(
    filename: str,
    content: str,
    document_id: str | None = None,
    knowledge_base_id: str = DEFAULT_KNOWLEDGE_BASE,
) -> dict:
    validate_filename(filename)
    if len(content.strip()) < 10:
        raise DocumentError("文档至少需要 10 个字符。")
    return {
        "id": document_id
        or hashlib.sha256(
            (
                filename
                if knowledge_base_id == DEFAULT_KNOWLEDGE_BASE
                else knowledge_base_id + ":" + filename.casefold()
            ).encode()
        ).hexdigest()[:16],
        "knowledge_base_id": knowledge_base_id,
        "filename": filename,
        "title": next(
            (line.lstrip("# ") for line in content.splitlines() if line.startswith("# ")),
            Path(filename).stem,
        ),
        "content": content,
        "version": hashlib.sha256(content.encode()).hexdigest()[:10],
    }


def document_chunks(document: dict) -> list[Citation]:
    chunks = []
    content, filename = document["content"], document["filename"]
    for part in re.split(r"(?m)^## ", content):
        if not part.strip():
            continue
        heading, _, body = part.partition("\n")
        if not filename.lower().endswith(".md") or not content.lstrip().startswith("#"):
            heading, body = "正文", part
        if not body.strip():
            continue
        # 360 characters + at most 80 heading characters fit BGE's 512-token budget.
        # Overlap preserves facts that straddle boundaries.
        for offset in range(0, len(body), 300):
            text = body[offset : offset + 360].strip()
            if not text:
                continue
            number = len(chunks)
            chunk_id = hashlib.sha256(
                f"{document['id']}:{document['version']}:{number}".encode()
            ).hexdigest()[:20]
            chunks.append(
                Citation(
                    chunk_id=chunk_id,
                    document_id=document["id"],
                    document_name=filename,
                    heading=heading.lstrip("# "),
                    text=text,
                    version=document["version"],
                    knowledge_base_id=document.get("knowledge_base_id", DEFAULT_KNOWLEDGE_BASE),
                )
            )
            if offset + 360 >= len(body):
                break
    return chunks


class DocumentRepository:
    def __init__(self, config: Settings):
        self.path, self.config = config.database_path, config
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.knowledge = KnowledgeRepository(self.path)
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS kb_documents (
                    id TEXT PRIMARY KEY, filename TEXT UNIQUE COLLATE NOCASE,
                    title TEXT NOT NULL, content TEXT NOT NULL, version TEXT NOT NULL,
                    indexed_version TEXT NOT NULL DEFAULT '', deleted INTEGER NOT NULL DEFAULT 0,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS kb_document_versions (
                    id INTEGER PRIMARY KEY, document_id TEXT NOT NULL, filename TEXT NOT NULL,
                    content TEXT NOT NULL, version TEXT NOT NULL, archived_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS kb_meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
                INSERT OR IGNORE INTO kb_meta VALUES ('revision', '0');
                INSERT OR IGNORE INTO kb_meta VALUES ('index_revision', '-1');
            """)
            # Preserve existing IDs/history, while allowing the same filename in different bases.
            db.execute("BEGIN IMMEDIATE")
            columns = {r[1] for r in db.execute("PRAGMA table_info(kb_documents)")}
            if "knowledge_base_id" not in columns:
                db.execute("""CREATE TABLE kb_documents_v2 (
                    id TEXT PRIMARY KEY, filename TEXT NOT NULL COLLATE NOCASE,
                    title TEXT NOT NULL, content TEXT NOT NULL, version TEXT NOT NULL,
                    indexed_version TEXT NOT NULL DEFAULT '', deleted INTEGER NOT NULL DEFAULT 0,
                    updated_at TEXT NOT NULL, knowledge_base_id TEXT NOT NULL DEFAULT 'default',
                    UNIQUE(knowledge_base_id, filename)
                )""")
                db.execute(
                    "INSERT INTO kb_documents_v2 SELECT "
                    "id,filename,title,content,version,'',deleted,updated_at,'default' "
                    "FROM kb_documents"
                )
                db.execute("DROP TABLE kb_documents")
                db.execute("ALTER TABLE kb_documents_v2 RENAME TO kb_documents")
                db.execute("UPDATE kb_meta SET value=CAST(value AS INTEGER)+1 WHERE key='revision'")

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA journal_mode=WAL")
        try:
            with db:
                yield db
        finally:
            db.close()

    def revision(self) -> int:
        with self.connect() as db:
            return int(db.execute("SELECT value FROM kb_meta WHERE key='revision'").fetchone()[0])

    def summary(self, document: dict) -> dict:
        return {k: v for k, v in document.items() if k not in {"content", "deleted"}} | {
            "characters": len(document["content"]),
            "chunks": len(document_chunks(document)),
            "index_status": "indexed"
            if document["indexed_version"] == document["version"]
            else "pending"
            if self.config.vector_enabled
            else "lexical",
        }

    def list(self, *, content=False, knowledge_base_id: str | None = None) -> list[dict]:
        with self.connect() as db:
            sql = "SELECT * FROM kb_documents WHERE deleted=0"
            args = ()
            if knowledge_base_id is not None:
                sql += " AND knowledge_base_id=?"
                args = (knowledge_base_id,)
            docs = [dict(r) for r in db.execute(sql + " ORDER BY filename", args)]
        return docs if content else [self.summary(doc) for doc in docs]

    def get(self, document_id: str) -> dict:
        with self.connect() as db:
            row = db.execute(
                "SELECT * FROM kb_documents WHERE id=? AND deleted=0", (document_id,)
            ).fetchone()
        if row is None:
            raise KeyError(document_id)
        return dict(row)

    def by_filename(self, filename: str, knowledge_base_id=DEFAULT_KNOWLEDGE_BASE) -> dict | None:
        with self.connect() as db:
            row = db.execute(
                "SELECT * FROM kb_documents WHERE filename=? AND knowledge_base_id=? AND deleted=0",
                (filename, knowledge_base_id),
            ).fetchone()
        return dict(row) if row else None

    def snapshot(self, *, include_revision=False):
        with self.connect() as db:
            db.execute("BEGIN")
            revision = int(
                db.execute("SELECT value FROM kb_meta WHERE key='revision'").fetchone()[0]
            )
            docs = [
                dict(row)
                for row in db.execute(
                    "SELECT * FROM kb_documents WHERE deleted=0 ORDER BY filename"
                )
            ]
        snapshot = ([self.summary(d) for d in docs], [c for d in docs for c in document_chunks(d)])
        return (*snapshot, revision) if include_revision else snapshot

    def save(self, doc: dict, *, indexed=False):
        knowledge_base_id = doc.get("knowledge_base_id", DEFAULT_KNOWLEDGE_BASE)
        self.knowledge.get(knowledge_base_id)
        now = datetime.now(UTC).isoformat()
        with self.connect() as db:
            old = db.execute("SELECT * FROM kb_documents WHERE id=?", (doc["id"],)).fetchone()
            if old:
                db.execute(
                    "INSERT INTO kb_document_versions "
                    "(document_id,filename,content,version,archived_at) VALUES (?,?,?,?,?)",
                    (old["id"], old["filename"], old["content"], old["version"], now),
                )
            db.execute(
                "INSERT INTO kb_documents "
                "(id,filename,title,content,version,indexed_version,deleted,updated_at,"
                "knowledge_base_id) "
                "VALUES (?,?,?,?,?,?,0,?,?) ON CONFLICT(id) DO UPDATE SET "
                "filename=excluded.filename,title=excluded.title,content=excluded.content,"
                "version=excluded.version,indexed_version=excluded.indexed_version,"
                "deleted=0,updated_at=excluded.updated_at,knowledge_base_id=excluded.knowledge_base_id",
                (
                    *[doc[k] for k in ("id", "filename", "title", "content", "version")],
                    doc["version"] if indexed else "",
                    now,
                    knowledge_base_id,
                ),
            )
            db.execute("UPDATE kb_meta SET value=CAST(value AS INTEGER)+1 WHERE key='revision'")
        return self.summary(self.get(doc["id"]))

    def delete(self, document_id: str):
        self.get(document_id)
        with self.connect() as db:
            db.execute("UPDATE kb_documents SET deleted=1 WHERE id=?", (document_id,))
            db.execute("UPDATE kb_meta SET value=CAST(value AS INTEGER)+1 WHERE key='revision'")

    def mark_indexed(self, revision: int):
        with self.connect() as db:
            current = int(
                db.execute("SELECT value FROM kb_meta WHERE key='revision'").fetchone()[0]
            )
            if current == revision:
                db.execute("UPDATE kb_documents SET indexed_version=version WHERE deleted=0")
                db.execute(
                    "UPDATE kb_meta SET value=? WHERE key='index_revision'", (str(revision),)
                )

    def delete_knowledge_base(self, knowledge_base_id: str):
        """Remove the aggregate atomically; retain historical messages and citation snapshots."""
        if knowledge_base_id == DEFAULT_KNOWLEDGE_BASE:
            raise ValueError("默认知识库作为入口保留，可以删除其中的文档。")
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            if not db.execute(
                "SELECT 1 FROM knowledge_bases WHERE id=?", (knowledge_base_id,)
            ).fetchone():
                raise KeyError(knowledge_base_id)
            if db.execute(
                "SELECT 1 FROM sessions WHERE knowledge_base_id=? AND status='active'",
                (knowledge_base_id,),
            ).fetchone():
                raise ValueError("此知识库有进行中的通话，请先结束通话。")
            db.execute(
                "DELETE FROM kb_document_versions WHERE document_id IN "
                "(SELECT id FROM kb_documents WHERE knowledge_base_id=?)",
                (knowledge_base_id,),
            )
            db.execute("DELETE FROM kb_documents WHERE knowledge_base_id=?", (knowledge_base_id,))
            db.execute("DELETE FROM knowledge_bases WHERE id=?", (knowledge_base_id,))
            db.execute(
                "UPDATE sessions SET status='ended',updated_at=? WHERE knowledge_base_id=?",
                (datetime.now(UTC).isoformat(), knowledge_base_id),
            )
            db.execute("UPDATE kb_meta SET value=CAST(value AS INTEGER)+1 WHERE key='revision'")
