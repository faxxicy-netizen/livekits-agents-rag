import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app.knowledge.retrieval import RetrievalRegistry
from app.knowledge.vector import VectorError
from app.main import create_app


def test_upload_edit_delete_and_no_demo_resurrection(config):
    app = create_app(config)
    with TestClient(app) as client:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("会议室.txt", "明川园区会议室需要提前三小时预约。".encode("gb18030"))},
        )
        assert response.status_code == 200
        doc = response.json()["document"]
        assert "三小时" in client.get("/api/documents/" + doc["id"]).json()["content"]
        client.put(
            "/api/documents/" + doc["id"],
            json={
                "content": "明川园区会议室需要提前五小时预约。",
                "expected_version": doc["version"],
            },
        ).raise_for_status()
        assert (
            client.put(
                "/api/documents/" + doc["id"],
                json={
                    "content": "明川园区会议室需要提前六小时预约。",
                    "expected_version": doc["version"],
                },
            ).status_code
            == 409
        )
        for item in client.get("/api/documents").json():
            client.delete("/api/documents/" + item["id"]).raise_for_status()
        assert client.get("/api/documents/" + doc["id"]).status_code == 404
    with TestClient(create_app(config)) as reopened:
        assert reopened.get("/api/documents").json() == []


def test_docx_paragraphs_tables_and_scanned_pdf_feedback(config):
    document = Document()
    document.add_paragraph("明川园区的访客预约流程如下，请提前登记。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "预约截止时间", "前一工作日十七点"
    stream = io.BytesIO()
    document.save(stream)
    with TestClient(create_app(config)) as client:
        response = client.post(
            "/api/documents/upload", files={"file": ("访客.docx", stream.getvalue())}
        )
        response.raise_for_status()
        content = client.get("/api/documents/" + response.json()["document"]["id"]).json()[
            "content"
        ]
        assert "提前登记" in content and "十七点" in content
        pdf = PdfWriter()
        pdf.add_blank_page(width=200, height=200)
        blank = io.BytesIO()
        pdf.write(blank)
        response = client.post(
            "/api/documents/upload", files={"file": ("扫描件.pdf", blank.getvalue())}
        )
        assert response.status_code == 400 and "OCR" in response.json()["detail"]


@pytest.mark.asyncio
async def test_long_lived_retriever_sees_updates_from_another_instance(config):
    reader = RetrievalRegistry(config)
    with TestClient(create_app(config)) as writer:
        doc = writer.post(
            "/api/documents",
            json={
                "filename": "访客管理.txt",
                "content": "明川园区访客通行暗号是星河青鸟，登记后领取访客证。",
            },
        ).json()["document"]
        results = await reader.search("访客通行暗号", "bm25")
        assert any("星河青鸟" in c.text for c in results.chunks)
        writer.put(
            "/api/documents/" + doc["id"],
            json={
                "content": "明川园区访客通行暗号是银色海豚，登记后领取访客证。",
            },
        ).raise_for_status()
        results = await reader.search("访客通行暗号", "bm25")
        assert any("银色海豚" in c.text for c in results.chunks)
        assert not any("星河青鸟" in c.text for c in results.chunks)
        writer.delete("/api/documents/" + doc["id"]).raise_for_status()
        assert not any(
            c.document_id == doc["id"] for c in (await reader.search("访客通行暗号", "bm25")).chunks
        )


def test_failed_vector_write_never_publishes_new_document(config, monkeypatch):
    app = create_app(config)
    manager = app.state.registry.manager

    async def fail(*args):
        raise VectorError("test vector outage")

    monkeypatch.setattr(manager.index, "upsert", fail)
    with TestClient(app) as client:
        config.vector_enabled = True  # Enable after lifespan; no external service is used.
        response = client.post(
            "/api/documents",
            json={
                "filename": "失败写入.txt",
                "content": "这个文档不能在向量索引失败时被当作上传成功。",
            },
        )
        assert response.status_code == 503
        assert manager.repository.by_filename("失败写入.txt") is None
